import os
import sys
import io
import openpyxl
from django.core.management.base import BaseCommand
from bookings.models import LotteryResult

class Command(BaseCommand):
    help = 'Imports lottery results from an Excel file into the LotteryResult table.'

    def add_arguments(self, parser):
        parser.add_argument('file_path', type=str, help='Path to the Excel file')
        parser.add_argument('--clear', action='store_true', help='Clear all existing lottery results before import')

    def handle(self, *args, **kwargs):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
        file_path = kwargs['file_path']
        clear_first = kwargs.get('clear', False)

        if not os.path.exists(file_path):
            print(f'[ERROR] File "{file_path}" does not exist.')
            return

        try:
            print(f'[INFO] Reading {file_path}...')
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            
            if clear_first:
                deleted_count = LotteryResult.objects.count()
                LotteryResult.objects.all().delete()
                print(f'[INFO] Cleared {deleted_count} existing lottery results.')
            
            # Read headers from first row
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            # Expected: العضوية, رقم العضوية, الاســــــم, النادى, الوصف, رقم الوحدة, الفترة, التليفون
            
            imported = 0
            updated = 0
            skipped = 0
            
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) < 7:
                    skipped += 1
                    continue
                
                # row: (العضوية, رقم العضوية, الاسم, النادى, الوصف, رقم الوحدة, الفترة, التليفون)
                _, mem_no_val, member_name, club_name, unit_desc, unit_no_val, period_val = row[:7]
                phone_val = row[7] if len(row) > 7 else ''
                
                membership_no = str(int(mem_no_val)) if mem_no_val is not None else ''
                member_name = str(member_name) if member_name else ''
                club_name = str(club_name) if club_name else ''
                unit_description = str(unit_desc) if unit_desc else ''
                unit_no = str(int(unit_no_val)) if unit_no_val is not None else ''
                period_no = int(period_val) if period_val is not None else 0
                phone = str(phone_val) if phone_val else ''
                # Clean phone
                if phone.endswith('.0'):
                    phone = phone[:-2]

                if not membership_no or not unit_no:
                    skipped += 1
                    continue

                _, created = LotteryResult.objects.update_or_create(
                    membership_no=membership_no,
                    club_name=club_name,
                    unit_no=unit_no,
                    period_no=period_no,
                    defaults={
                        'member_name': member_name,
                        'unit_description': unit_description,
                        'phone': phone,
                    }
                )
                if created:
                    imported += 1
                else:
                    updated += 1

            wb.close()
            total_in_db = LotteryResult.objects.count()
            print(f'[SUCCESS] Lottery Import Complete:')
            print(f'  - New records: {imported}')
            print(f'  - Updated records: {updated}')
            print(f'  - Skipped (empty data): {skipped}')
            print(f'  - Total in database: {total_in_db}')
            
            # Automatically sync imported results to bookings
            print('\n[INFO] Running automatic synchronization to Bookings...')
            from bookings.services import sync_lottery_results_to_bookings
            sync_lottery_results_to_bookings()

        except Exception as e:
            print(f'[ERROR] Import failed: {str(e)}')
            import traceback
            traceback.print_exc()
