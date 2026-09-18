import openpyxl
import json

try:
    wb = openpyxl.load_workbook(r"d:\housing\الاسعار.xlsx", data_only=True)
    res = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = []
        for row in ws.iter_rows(min_row=1, max_row=25, values_only=True):
            # Convert None to ""
            cleaned_row = [str(cell) if cell is not None else "" for cell in row]
            # Only append if the row is not completely empty
            if any(cleaned_row):
                data.append(cleaned_row)
        res[sheet_name] = data

    with open(r"d:\housing\backend\excel_data.json", "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=2)
except Exception as e:
    print(f"Error: {e}")
