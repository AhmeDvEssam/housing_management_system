"""فحص عمود التصنيف في Data.xlsx"""
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from collections import Counter

wb = openpyxl.load_workbook(r'd:\housing\Data.xlsx', read_only=True)
ws = wb.active

headers = [c.value for c in ws[1]]
print('Headers:', headers)

# Find classification column index
class_idx = None
for i, h in enumerate(headers):
    if h and 'تصنيف' in str(h):
        class_idx = i
        break

print(f'Classification column index: {class_idx}')

classifications = []
rows_sample = []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    val = row[class_idx] if class_idx is not None and class_idx < len(row) else None
    classifications.append(val)
    if i <= 15:
        rows_sample.append(row)
        print(f'  Row {i}: Club={row[0]} | Building={row[1]} | Unit={row[2]} | Type={row[3]} | Classification={row[class_idx] if class_idx is not None else "N/A"} | Floor={row[5] if len(row) > 5 else "N/A"}')

c = Counter(classifications)
print(f'\n=== توزيع التصنيفات في Excel ===')
for k, v in c.most_common():
    print(f'  {repr(k)}: {v}')
print(f'\nTotal rows: {len(classifications)}')
wb.close()
