from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Iterable

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def parse_date(value: str | None, default: date | None = None) -> date | None:
    if not value:
        return default
    try:
        return datetime.strptime(value, "%d-%m-%Y").date()
    except ValueError as exc:
        raise ValueError("صيغة التاريخ يجب أن تكون DD-MM-YYYY") from exc


def resolve_date_range(
    *,
    start_date_value: str | None,
    end_date_value: str | None,
    days_value: str | None,
    default_days: int = 7,
) -> tuple[date, date]:
    start_date = parse_date(start_date_value, timezone.localdate())
    if not start_date:
        start_date = timezone.localdate()
    end_date = parse_date(end_date_value)

    if not end_date:
        days = default_days
        if days_value:
            try:
                days = int(days_value)
            except ValueError:
                raise ValueError("قيمة عدد الأيام غير صالحة")
        if days < 1:
            days = 1
        if days > 31:
            days = 31
        end_date = start_date + timedelta(days=days)

    if end_date <= start_date:
        raise ValueError("تاريخ النهاية يجب أن يكون بعد تاريخ البداية")

    return start_date, end_date


def should_export_xlsx(format_value: str | None) -> bool:
    return bool(format_value and format_value.lower() == "xlsx")


def build_workbook(
    *,
    title: str,
    filter_text: str,
    headers: Iterable[str],
    rows: Iterable[Iterable[object]],
    sheet_name: str,
    filename: str,
) -> HttpResponse:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    worksheet.append([title])
    worksheet.append([filter_text])
    worksheet.append([])
    worksheet.append(list(headers))

    for row in rows:
        worksheet.append(list(row))

    for column_cells in worksheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = (
            min(max(length + 2, 10), 40)
        )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type=
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename=\"{filename}\""
    return response
